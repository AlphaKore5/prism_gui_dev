"""BOQ-CHANGES-TODO.md item 30 — the "Prism Agent" write-up path.

Every other writer agent (Claude, ChatGPT, Grok...) is a real tool in a
browser, which can export its own answer as a DOCX/PDF because that tool
has its own file-making feature built in. A plain OpenRouter chat
completion has none of that — it only ever returns text. So this module
is the other half of item 30b: ask the model for a well-structured
answer (headings, a pipe-table of BOQ lines, same shape the write-up
prompt has always asked every agent for), then build the actual
DOCX/CSV/PDF file from that text here, in code, the same way boq_price.py
already builds a document from structured data rather than asking an AI
to produce a file directly.

Deliberately NOT routed through core.boq_price's own Boq/BoqItem/
write_boq_* pipeline: that module's whole premise, stated in its own
docstring, is "no number of record is ever produced by an AI" — it is
the PRICING stage's own deterministic, rate-library-matched pipeline.
The write-up stage has never worked that way; it has always handed back
free-form text built around measured figures that were fixed before the
prompt was written (see dialog.py's _write/_send_followup, which embed
`self.summary` as plain text). This module keeps that same shape — it
is a renderer for the write-up's own text, not a second pricing engine.
"""
from __future__ import annotations

import csv
import html as _html
import os
import re

import core_bridge as CB
from addons.boq import agents as boq_agents


def model_for(agent: str, *, followup: bool) -> str:
    """Item 30e's two-tier choice: the heavier generation model, or the
    cheaper same-family sibling for a follow-up/edit."""
    table = (boq_agents.PRISM_AGENT_FOLLOWUP_MODELS if followup
            else boq_agents.PRISM_AGENT_MODELS)
    model = table.get(agent)
    if not model:
        raise ValueError(f"{agent!r} is not a Prism Agent.")
    return model


def has_vision(agent: str) -> bool:
    """item 32: can this agent actually look at a rendered drawing at
    all. dialog.py checks this BEFORE calling write() with attachments,
    so it can tell the user plainly (Prism Agent Deep has none, at any
    tier) rather than silently sending the drawing and having it ignored
    with no explanation."""
    return bool(boq_agents.PRISM_AGENT_HAS_VISION.get(agent))


def write(agent: str, prompt: str, api_key: str, *, followup: bool = False,
         attachments: list[dict] | None = None) -> str:
    """One Prism Agent completion. Raises RuntimeError with the real
    reason on failure (no key, a rejected key, a network error, ...) --
    caught and shown to the user by whatever calls this off the GUI
    thread (addons/boq/workers.py:OpenRouterWriteWorker).

    `attachments`: rendered drawings (cad_render.prepare_for_ai()'s own
    shape) to show the model alongside the text, item 32 -- silently
    dropped here if `agent` has no vision capability at all (has_vision()
    already told the caller that before this was ever called; this is
    the last line of defence, not the place a person first learns it)."""
    model = model_for(agent, followup=followup)
    if attachments and not has_vision(agent):
        attachments = None
    return CB.openrouter.openrouter_chat(api_key, model, prompt, attachments=attachments)


# ── a small markdown model -- just enough for a BOQ write-up's own shape ──
# Headings, paragraphs, bullet/numbered lists, pipe tables. Not a general
# markdown engine -- the write-up prompt has always asked for exactly
# these four things (see prompts.py's formatting_prompt), so parsing only
# what it actually asks for is a smaller, more honest surface than a
# general parser that silently mishandles the rest.

def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def _is_table_sep(line: str) -> bool:
    s = line.strip()
    if not (s.startswith("|") and s.endswith("|")):
        return False
    cells = [c.strip() for c in s.strip("|").split("|")]
    # each cell is dashes, optionally with a leading/trailing ':' for
    # alignment ("---", ":---", "---:", ":---:") -- a bare "|" stripped of
    # its own pipes leaves "-" characters PLUS more "|"s between cells,
    # which the old `all(c in "-: " for c in s)` check (scanning the whole
    # row as one string) wrongly failed on, since "|" itself was never in
    # that allowed set. Splitting into cells first is what fixes it.
    return bool(cells) and all(c and set(c) <= set("-:") for c in cells)


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_markdown(text: str) -> list[dict]:
    """[{'type': 'heading'|'para'|'bullet'|'table', ...}, ...] in document
    order. `heading`: level, text. `para`: text. `bullet`: items (a list
    of strings). `table`: headers, rows (list of list of strings)."""
    lines = (text or "").replace("\r\n", "\n").split("\n")
    blocks: list[dict] = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)),
                           "text": m.group(2).strip()})
            i += 1
            continue
        if _is_table_row(line) and i + 1 < n and _is_table_sep(lines[i + 1]):
            headers = _split_row(stripped)
            rows = []
            i += 2
            while i < n and _is_table_row(lines[i]):
                rows.append(_split_row(lines[i]))
                i += 1
            blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue
        if re.match(r"^([-*]|\d+\.)\s", stripped):
            items = []
            while i < n and re.match(r"^([-*]|\d+\.)\s", lines[i].strip()):
                items.append(re.sub(r"^([-*]|\d+\.)\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "bullet", "items": items})
            continue
        para = [stripped]
        i += 1
        while (i < n and lines[i].strip()
              and not re.match(r"^#{1,4}\s", lines[i].strip())
              and not _is_table_row(lines[i])
              and not re.match(r"^([-*]|\d+\.)\s", lines[i].strip())):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"type": "para", "text": " ".join(para)})
    return blocks


def markdown_to_docx(text: str, path: str) -> str:
    import docx
    blocks = parse_markdown(text)
    doc = docx.Document()
    for b in blocks:
        if b["type"] == "heading":
            doc.add_heading(b["text"], level=min(b["level"], 4))
        elif b["type"] == "para":
            doc.add_paragraph(b["text"])
        elif b["type"] == "bullet":
            for item in b["items"]:
                doc.add_paragraph(item, style="List Bullet")
        elif b["type"] == "table" and b["headers"]:
            cols = len(b["headers"])
            t = doc.add_table(rows=1, cols=cols)
            try:
                t.style = "Light Grid Accent 1"
            except KeyError:                                 # template-dependent
                pass
            for c, h in enumerate(b["headers"]):
                t.rows[0].cells[c].text = h
            for row in b["rows"]:
                cells = t.add_row().cells
                for c, val in enumerate(row[:cols]):
                    cells[c].text = val
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    doc.save(path)
    return path


def markdown_to_csv(text: str, path: str) -> str:
    """The BIGGEST table in the answer (most rows) -- almost always the
    BOQ item schedule itself, not a small preamble table. No table found
    at all -> the plain text, one line per row, so a CSV is still
    produced rather than silently nothing."""
    blocks = parse_markdown(text)
    tables = [b for b in blocks if b["type"] == "table"]
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if not tables:
            w.writerow(["text"])
            for line in (text or "").splitlines():
                if line.strip():
                    w.writerow([line.strip()])
            return path
        table = max(tables, key=lambda t: len(t["rows"]))
        w.writerow(table["headers"])
        for row in table["rows"]:
            w.writerow(row)
    return path


def _markdown_to_html(text: str) -> str:
    blocks = parse_markdown(text)
    parts = [
        "<html><head><meta charset='utf-8'><style>",
        "body{font-family:Arial,Helvetica,sans-serif;font-size:11px;padding:16px}",
        "table{border-collapse:collapse;width:100%;margin:10px 0}",
        "th,td{border:1px solid #999;padding:4px 6px;text-align:left;"
        "vertical-align:top}",
        "th{background:#eee}",
        "h1{font-size:18px}h2{font-size:15px}h3{font-size:13px}h4{font-size:12px}",
        "</style></head><body>",
    ]
    for b in blocks:
        if b["type"] == "heading":
            lvl = min(b["level"], 4)
            parts.append(f"<h{lvl}>{_html.escape(b['text'])}</h{lvl}>")
        elif b["type"] == "para":
            parts.append(f"<p>{_html.escape(b['text'])}</p>")
        elif b["type"] == "bullet":
            items = "".join(f"<li>{_html.escape(i)}</li>" for i in b["items"])
            parts.append(f"<ul>{items}</ul>")
        elif b["type"] == "table":
            head = "".join(f"<th>{_html.escape(h)}</th>" for h in b["headers"])
            rows = "".join(
                "<tr>" + "".join(f"<td>{_html.escape(c)}</td>" for c in row) + "</tr>"
                for row in b["rows"])
            parts.append(f"<table><tr>{head}</tr>{rows}</table>")
    parts.append("</body></html>")
    return "".join(parts)


_PDF_FOOTER = (
    "<div style='font-size:8px;width:100%;text-align:center;color:#888;'>"
    "<span class='pageNumber'></span> / <span class='totalPages'></span></div>")


def markdown_to_pdf(text: str, path: str) -> str:
    """Same technique core.boq_price.write_boq_pdf already uses: render an
    HTML page with the bundled Chromium (Playwright) and print it."""
    html_str = _markdown_to_html(text)
    try:
        from playwright.sync_api import sync_playwright
    except Exception as e:                                  # noqa: BLE001
        raise RuntimeError(
            "Exporting a PDF needs the bundled browser engine "
            "(Playwright/Chromium), which isn't available here. DOCX and "
            "CSV still work without it.") from e
    from core import browser
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with sync_playwright() as pw:
        b = browser.launch_chromium(
            pw, args=["--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage"])
        try:
            page = b.new_page()
            page.set_content(html_str, wait_until="load")
            page.pdf(path=path, format="A4", print_background=True,
                    margin={"top": "14mm", "bottom": "16mm",
                           "left": "12mm", "right": "12mm"},
                    display_header_footer=True,
                    header_template="<span></span>",
                    footer_template=_PDF_FOOTER)
        finally:
            b.close()
    return path


FORMAT_WRITERS = {
    "docx": markdown_to_docx,
    "csv": markdown_to_csv,
    "pdf": markdown_to_pdf,
}


def write_file(text: str, fmt: str, path: str) -> str:
    writer = FORMAT_WRITERS.get((fmt or "").lower())
    if writer is None:
        raise ValueError(f"No writer for format {fmt!r}.")
    return writer(text, path)
